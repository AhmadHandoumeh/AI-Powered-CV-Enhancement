from docx import Document


def extract_docx_text(file):
    file.seek(0)
    document = Document(file)

    paragraphs = [
        p.text.strip()
        for p in document.paragraphs
        if p.text.strip()
    ]

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())

    return "\n".join(paragraphs)